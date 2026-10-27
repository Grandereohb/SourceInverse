"""Concise editable Chapter4.4 report, backed by all frozen public releases."""

import argparse
import hashlib
from pathlib import Path

import pandas as pd

from run_focused_chapter import ROOT,STUDY,read,save

OUT=STUDY/'metec_comparison_20261027_endtime_q16'
METHODS=('ia_vpdi','joint_nls_trf','standard_gbest_pso','matched_joint_adam')

LABELS=['IA-VPDI','Joint NLS','标准PSO','联合Adam']


def add_chapter(document,page_break_before=False):
    from docx.shared import Cm,Pt
    from build_paper_results_summary import paragraph,heading,table
    summary=read(OUT/'analysis/summary.json')
    if not summary['final'] or summary['complete_events']!=17:
        raise ValueError('Complete METEC cohort required')
    for path,digest in read(OUT/'analysis/sources.json').items():
        if hashlib.sha256((ROOT/path).read_bytes()).hexdigest()!=digest:
            raise ValueError('Fit source changed after evaluation: '+path)
    for name in ['all_releases','first_release']:
        check=read(OUT/'figures'/(name+'.collisions.json'))
        if any(str(f.get('severity',f.get('level',''))).upper()=='FAIL' for f in check.get('findings',[])):
            raise ValueError('Figure QA failure: '+name)
    frame=pd.read_csv(OUT/'analysis/case_metrics.csv',dtype={'event_id':str})
    raw=pd.read_csv(OUT/'analysis/replicate_metrics.csv',dtype={'event_id':str})
    failed=raw[raw.replay_passed.astype('string').str.lower().eq('false')]
    stats=pd.read_csv(OUT/'analysis/method_summary.csv').set_index('method')
    def p(text,bold=False):
        item=paragraph(document,text,bold,size=10)
        item.paragraph_format.space_after=Pt(5)
        item.paragraph_format.line_spacing=1.05
    def grid(headers,rows,widths):
        table(document,headers,rows,widths)
    def image(name,caption,width=17):
        document.add_picture(str(OUT/'figures'/(name+'.png')),width=Cm(width))
        paragraph(document,caption,size=9)
    def pair(method,metric,mult=1,digits=2):
        row=stats.loc[method]
        return f'{mult*row[metric+"_median"]:.{digits}f} / {mult*row[metric+"_mean"]:.{digits}f}'
    heading(document,'4.4 METEC公开实测定位验证')
    document.paragraphs[-1].paragraph_format.page_break_before=page_break_before
    p('目的：以独立释放坐标与计量排放量为真值，检验方法在公开多站实测数据上的定位与排放量重建能力。'
      '全部符合条件的17次释放均纳入统计。')
    p('数据为Daniels等公开的一周甲烷样本，UTC 2022年5月9日至15日，8个浓度站、3处风观测。'
      '17次释放涉及5个设备组、11处记录坐标，均为单源且互不重叠；不是17个独立场地。'
      '真实甲烷流率约0.54至6.11 kg/h，释放点坐标和流率仅用于事后评价。')
    p('四种方法共用三维点观测输运、地面反射、源高、全部有效观测、搜索边界与两个初值。'
      '源高采用设备组记录的2或4.5米，站高2.4米；所以这是源高已知条件下的定位。'
      'IA自动选择1、3或5个排放参数，其余方法固定5个。PSO每例3种子取指标中位数；预算不同，不称等计算量比较。')
    rows=[[label,pair(m,'localization_error_m'),pair(m,'q_relative_error',100,1),pair(m,'rmse_ppm',digits=3)]
        for m,label in zip(METHODS,LABELS)]
    grid(['方法','定位误差 米','源强误差 %','浓度误差 ppm'],rows,[3.0,4.3,4.3,5.4])
    p('表内均为中位数 / 平均数，分母为17次释放。源强误差比较重建时变源强的时间均值与计量均值，'
      '不代表真实逐分钟排放曲线已知。浓度误差针对扣除背景后的全部有效观测。')
    p(f'17例结果完整；数值核查{summary["replay_passes"]}/{summary["replay_total"]}项通过，'
      f'{failed.event_id.nunique()}个事件待精修，详见末页。源强均值含异常大值，尚非全部验收通过的定稿。',True)
    difference=summary['ia_minus_nls_location_m']
    p(f'IA定位改善超过1米的有{difference["ia_better_by_more_than_1m"]}例，NLS改善超过1米的有'
      f'{difference["nls_better_by_more_than_1m"]}例，其余{difference["within_1m"]}例相差不超过1米。1米是描述口径，不是显著性标准。',True)
    image('all_releases','图1 全部17次释放的逐例误差。横轴按时间排序，左为定位，右为平均源强；'
          '纵轴采用保留零值的对称对数刻度。PSO为例内3种子中位数，灰带为最小至最大范围，不是置信区间。',16)

    heading(document,'4.4 代表事件与数据处理')
    document.paragraphs[-1].paragraph_format.page_break_before=True
    first=read(OUT/'protocol.json')['event_ids'][0]
    image('first_release',f'图2 预先按时间选取的首个完整事件{first}。左为站点、真源与估计位置，'
        '右为观测峰值最高站的浓度增量。PSO展示固定种子1729；点线重合表示估计接近。')
    p('输入处理：重复的同站同分钟记录取均值；浓度缺测不补零；各风站先转为矢量再求均值，'
      '全站风缺测仅允许不超过5分钟的短缺口插值。背景取释放前30分钟各站20分位；'
      '某站事前全缺时，用前三小时内最近30个有效分钟估计并记录。')
    p('物理口径：源强单位kg/s，输出换算为kg/h；浓度使用作者固定转换1.524×10⁶ ppm/(kg/m³)，'
      '并非实测温压。沿风与横风扩散随风向旋转，竖直扩散取横风扩散系数，地面镜像反射；'
      '分钟值按作者定义对应截至时间戳的前一分钟，使用积分计算；已知释放起止保留秒精度，'
      '区间内排放量未知、区间外为零，不测试泄漏自动检出。')
    p('位置边界仅取站点包围框向外扩25米，扩散范围事先固定；初值来自共同粗网格对观测的拟合。'
      '不把真实释放位置、设备中心或计量流率作为初值或搜索提示；源高是共同已知条件。')
    p(f'数值核查：在所有{summary["method_runs"]}次已保存方法运行处，将时间/释放积分由16/16阶提高为32/32阶，'
      f'历史由20分钟延长至40分钟，仅回放、不重拟合。原1%相对预测差门槛通过'
      f'{summary["replay_passes"]}/{summary["replay_total"]}项，最大差{100*summary["replay_max_relative_l2"]:.3f}%。'
      '未通过项保留并在逐例CSV标记；回放稳定不等于物理模型完全正确或位置唯一。')
    counts=raw[raw.method=='ia_vpdi'].q_basis_count.value_counts().sort_index()
    p('IA所选排放参数数量：'+ '，'.join(f'{int(k)}个参数{int(v)}例' for k,v in counts.items())+'。'
      '复杂度选择使用全部观测，不留出站点或时间块。模型可表达时变排放，但计量真值只提供事件均值。')
    convergence=[]
    for method,label in zip(METHODS[:2],LABELS[:2]):
        group=raw[raw.method==method]
        count=int(group.converged.astype('string').str.lower().eq('true').sum())
        convergence.append(f'{label}优化器报告收敛{count}/17例')
    p('，'.join(convergence)+'。未收敛案例也保留；PSO与Adam按固定预算停止，不等同于已验证收敛。')
    p('资料出处：Daniels等2024，DOI 10.1525/elementa.2023.00110；公开代码与样本wsdaniels/DLQ，'
      '固定提交d289f32264d2d32045a1d486670af9a6dd136c80。坐标测量精度未给出，不按小数位数解释为厘米精度。')
    p('与原Chilbolton结果的关系：原8窗口及两种源高结果继续保留；METEC补充独立固定站网络、'
      '更多释放位置和较长事件，不能将两套不同观测模型的数据混为一个总体或直接比较绝对误差。')

    heading(document,'4.4 全部事件与待精修项')
    document.paragraphs[-1].paragraph_format.page_break_before=True
    wide=frame.pivot(index='event_id',columns='method',values='localization_error_m')
    ids=read(OUT/'protocol.json')['event_ids']
    rows=[[str(j+1)+' '+k,*[f'{wide.loc[k,m]:.2f}' for m in METHODS]] for j,k in enumerate(ids)]
    grid(['编号 释放ID',*LABELS],rows,[4.6,3.0,3.3,3.1,3.0])
    p('单位均为米。表内保留全部17次释放；PSO为三次技术重复的定位误差中位数，重复次数不扩大样本量。')
    p('源强均值受极大估计值影响。例如20220509006的计量均值约2.73 kg/h，PSO和Adam部分结果约61184 kg/h。'
      '异常保留，暂不作为可靠排放量；定位和源强重建分开评价。')
    grid(['待精修事件','超限运行数','最大预测变化 %'],[[event,len(group),f'{100*group.replay_relative_l2.max():.2f}']
        for event,group in failed.groupby('event_id')],[6.5,4.5,6.0])
    p('后续优先定向核对这4个事件的积分精度，以及20220509006的源强初始化、响应尺度和局部解；'
      '不需要为此另开一轮大规模场景实验。单位已通过独立积分核对，但源强异常的原因尚未完成归因。')
    ia,nls=stats.loc['ia_vpdi'],stats.loc['joint_nls_trf']
    p(f'本样本中IA与NLS定位中位误差约{ia.localization_error_m_median:.2f}和{nls.localization_error_m_median:.2f}米，'
      f'平均误差约{ia.localization_error_m_mean:.2f}和{nls.localization_error_m_mean:.2f}米。'
      '数据提供了独立实测定位对照，当前结果不支持IA具有明确定位优势；源强异常与数值警告需在论文定稿前处理。',True)


def build(output):
    from docx import Document
    from docx.shared import Cm,Pt,RGBColor
    from docx.oxml.ns import qn
    document=Document()
    section=document.sections[0]
    section.page_width,section.page_height=Cm(21),Cm(29.7)
    section.top_margin,section.bottom_margin=Cm(1.5),Cm(1.6)
    section.left_margin,section.right_margin=Cm(1.8),Cm(1.8)
    for name in ['Normal','Title','Heading 1']:
        style=document.styles[name]
        style.font.name='Arial'
        style.font.color.rgb=RGBColor(0,0,0)
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
        properties=style._element.find(qn('w:pPr'))
        if properties is not None:
            for border in list(properties.findall(qn('w:pBdr'))):
                properties.remove(border)
    title=document.add_paragraph('METEC公开实测验证报告',style='Title')
    title.paragraph_format.space_after=Pt(5)
    add_chapter(document)
    output.parent.mkdir(parents=True,exist_ok=True)
    document.save(output)
    save(output.with_suffix('.sources.json'),read(OUT/'analysis/sources.json'))
    print(output)


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output',type=Path,default=OUT/'METEC公开实测验证报告.docx')
    build(parser.parse_args().output)
