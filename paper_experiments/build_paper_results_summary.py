"""Create a short editable report from the complete writing-data packet."""

import argparse
import csv
from datetime import date
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]/'result/paper_experiments_20261017'
PACKET = ROOT/'manuscript_results_v1'
METHODS = {'ia_vpdi': 'IA-VPDI', 'matched_joint_adam': '联合 Adam',
           'shared_response_nnls_powell': 'NNLS-Powell'}


def rows(name, packet=None):
    with ((packet or PACKET)/name).open(encoding='utf-8-sig', newline='') as stream:
        return list(csv.DictReader(stream))


def font(run, size=10.5, bold=False):
    run.font.name = 'Arial'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')


def paragraph(document, text, bold=False, size=10.5):
    item = document.add_paragraph()
    item.paragraph_format.space_after = Pt(6)
    item.paragraph_format.line_spacing = 1.15
    font(item.add_run(text), size, bold)
    return item


def heading(document, text):
    item = document.add_heading(text, level=1)
    item.paragraph_format.space_before = Pt(5)
    item.paragraph_format.space_after = Pt(8)
    for run in item.runs:
        font(run, 15, True)


def table(document, headers, data, widths):
    item = document.add_table(rows=1, cols=len(headers))
    item.alignment = WD_TABLE_ALIGNMENT.CENTER
    item.autofit = False
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        node = OxmlElement('w:'+edge)
        for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            node.set(qn('w:'+key), value)
        borders.append(node)
    item._tbl.tblPr.append(borders)
    for column, width in zip(item.columns, widths):
        column.width = Cm(width)
    for i, values in enumerate([headers, *data]):
        cells = item.rows[0].cells if i == 0 else item.add_row().cells
        for j, (cell, value) in enumerate(zip(cells, values)):
            cell.width = Cm(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            properties = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side in ('top', 'left', 'bottom', 'right'):
                node = OxmlElement('w:'+side)
                node.set(qn('w:w'), '85')
                node.set(qn('w:type'), 'dxa')
                margins.append(node)
            properties.append(margins)
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'DCE6F1' if i == 0 else ('F5F7F9' if i % 2 == 0 else 'FFFFFF'))
            properties.append(shade)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            font(p.add_run(str(value)), 9, i == 0)
        row_properties = item.rows[i]._tr.get_or_add_trPr()
        row_properties.append(OxmlElement('w:cantSplit'))
        if i == 0:
            row_properties.append(OxmlElement('w:tblHeader'))
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def n(value, digits=2):
    return f'{float(value):.{digits}f}'


def build(packet=PACKET, report_date=None):
    manifest = json.loads((packet/'data_manifest.json').read_text(encoding='utf-8'))
    planned = manifest['supplement_planned']
    if manifest['main_completed'] != 120 or manifest['supplement_completed'] != planned:
        raise ValueError('Complete main and all planned supplementary cases before reporting')
    per_scenario = planned//2
    report_date = date.fromisoformat(report_date) if report_date else date.today()
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.7), Cm(2.0)
    section.footer_distance = Cm(.8)
    section.left_margin, section.right_margin = Cm(1.8), Cm(1.8)
    normal = document.styles['Normal']
    normal.font.name, normal.font.size = 'Arial', Pt(10.5)
    normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    title_properties = document.styles['Title']._element.find(qn('w:pPr'))
    if title_properties is not None:
        for border in list(title_properties.findall(qn('w:pBdr'))):
            title_properties.remove(border)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(footer.add_run('论文实验结果更新汇总  ·  '), 9)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)

    title = document.add_paragraph(style='Title')
    font(title.add_run('论文实验结果更新汇总'), 21, True)
    paragraph(document, report_date.strftime('%Y年%m月%d日')+'  |  供论文写作前核对', size=10)
    paragraph(document, '本次完成两个补充场景、一次定位区域校准，以及同公开释放事件的文献结果对照。名义条件下三种求解方法的定位中位误差处于相近量级；校准后的验证集定位区域包含率达到93.75%。以下按实验章节汇总方法、数据及写作依据。')
    heading(document, '4.1 合成案例中的定位与源强恢复')
    paragraph(document, '目的与方法：在已知真源的合成场中比较 IA-VPDI、联合 Adam 与 NNLS-Powell。三者使用相同前向响应、观测和源强基函数；真源坐标由模拟输入给定，定位误差为估计与真值的平面欧氏距离。共120例，含96例名义场景及24例风场或释放形状失配。')
    main = [r for r in rows('main_group_summary.csv',packet) if r['cohort'] == 'nominal']
    table(document, ['网络及案例数', '方法', '定位中位数 m', 'Q相对L2中位数'],
          [[('理想五站' if r['network']=='ideal5' else '实际六站布局')+'  n=48', METHODS[r['method']],
            n(r['localization_error_m_median']), n(100*float(r['q_relative_l2_median']))+'%'] for r in main],
          [4.3, 3.5, 4.3, 5.3])
    paragraph(document, '实际六站指沿用真实网络的站点几何布局，浓度与真源仍为合成数据。120例中严格定位最优计数：IA 25例、联合 Adam 47例、NNLS-Powell 48例；不据胜负删减主样本。固定预算结束不等于所有优化器均已收敛。', size=10)
    paragraph(document, '主要结果：理想五站的定位中位误差相近；实际六站布局下各方法误差均明显增大。24例失配结果单独保存在逐例及分组表中，不与96例名义数据混合解释。', bold=True)

    document.add_page_break()
    heading(document, '4.1 两个新增场景')
    paragraph(document, f'目的：检验站网几何与释放过程改变后的表现。五站环绕场景采用恒风与脉冲释放；六站走廊场景采用转风与双峰释放。每场景{per_scenario}例，共{planned}例，全部保留。扩样决定在初期结果可见后作出，新增种子在求解前固定；多个条件同时改变，不作单因素解释。')
    supplement = rows('supplement_group_summary.csv',packet)
    table(document, [f'场景 n={per_scenario}', '方法', '定位中位数 m', 'Q相对L2中位数'],
          [[('五站环绕脉冲' if r['network']=='surround5_pulse' else '六站走廊双峰'), METHODS[r['method']],
            n(r['localization_error_m_median']), n(100*float(r['q_relative_l2_median']))+'%'] for r in supplement],
          [4.3,3.5,4.3,5.3])
    heading(document, '4.2 定位可靠性与区域校准')
    paragraph(document, '目的：检验消除源强与扩散局部补偿后，定位区域是否更贴近真实不确定性。96例名义数据中，固定Q和D、消除Q、同时消除Q和D的未校准局部椭圆包含真源比例依次为16/96、41/96、80/96。')
    calibration = rows('calibration_coverage_summary.csv',packet)
    selected = [r for r in calibration if r['partition']=='validation']
    table(document, ['验证分组', '未校准包含率', '校准后包含率', '长半轴中位数 m'],
          [[{'all':'全部 n=48','ideal5':'理想五站 n=24','jjj6':'实际六站 n=24'}[r['network']],
            n(100*float(r['raw_coverage']))+'%',n(100*float(r['calibrated_coverage']))+'%',
            n(r['raw_major_median_m'],1)+' → '+n(r['calibrated_major_median_m'],1)] for r in selected],
          [4.3,3.5,4.3,5.3])
    paragraph(document, '一次校准：48例校准、另48例验证，按校准集马氏距离分位数将半轴统一乘3.343。验证包含率由39/48升至45/48，代价是区域扩大；不改变点定位，也不在源反演中留出站点或时段。既有模拟的回顾性分组尚不能证明所有场景均有95%覆盖保证。停止追加调参。', size=10)

    document.add_page_break()
    heading(document, '4.3 公开控制释放数据验证')
    paragraph(document, '数据与方法：Chilbolton公开甲烷控制释放数据，8个单源事件、2个物理源位置、7条光路。独立物理适配器处理垂直扩散、地面反射、光路平均及ppm换算。源高记录存在差异，0.3m和1.6m均保留，不择优选取。')
    physical = [r for r in rows('controlled_release_summary.csv',packet) if r['method']!='transient_hat5']
    labels = {'stationary_constant_q':'稳态恒Q','transient_constant_q':'瞬态恒Q','transient_hat5':'瞬态hat5'}
    table(document, ['源高 m', '模型', '8例定位中位数 m', '6例Q相对误差中位数'],
          [[r['height_m'],labels[r['method']],n(r['localization_median_m']),
            n(100*float(r['q_relative_error_median']))+'%'] for r in physical], [2.4,4.4,4.8,5.8])
    paragraph(document, '同事件文献对照：Vänskä等2025年 Inverse Problems 实验Case2表2同时报告rel06、rel07的2D/3D Bayesian state estimation结果。下表为同释放编号与同单位的已发表结果参照，不是完全同预处理的算法复现。')
    comparison = rows('controlled_release_same_event_comparison.csv',packet)
    data = []
    for r in comparison:
        method = {'published_2D_BSE':'文献2D BSE','published_3D_BSE':'文献3D BSE'}.get(r['method'],
                  '本文瞬态恒Q '+r['height_m']+'m')
        data.append([r['release'],method,n(r['localization_error_m']),n(r['estimated_q_kg_h'])])
    table(document, ['释放事件', '方法', '定位误差 m', '估计Q kg/h'],data,[2.4,6.8,4.0,4.2])
    paragraph(document, '文献参考Q为1.35/1.50kg/h（rel06/07），不替代缺失的独立称重值。rel07本文定位误差小于所列文献值，但Q明显偏高，须分开评价定位与源强。完整hat5等48组拟合结果均保存在CSV中。', size=10)

    document.add_page_break()
    heading(document, '写作结论与数据索引')
    paragraph(document, '4.1可写：相同输运条件下的求解比较，以及站网几何和模型失配对定位的影响。当前证据不支持IA整体精度第一；不能把IA获胜案例组成的子集写成总体评估。')
    paragraph(document, '4.2可写：消除局部补偿提高了定位区域对真源的描述能力；一次案例级尺度校准使独立分组的经验包含率达到93.75%。同时报告区域尺寸，不能把覆盖率提高写成点定位更准确。')
    paragraph(document, '4.3可写：在公开实测数据上达到米级定位，并与已有论文进行同事件数值对照。定位与源强结果分开讨论；源强偏差、时窗处理与源先验差异是解释结果的重点。')
    heading(document, '绘图与复核数据')
    paragraph(document, f'统一目录：result/paper_experiments_20261017/{packet.name}。CSV保留逐例结果，不仅保存均值；data_manifest.json记录原始输入、结果、优化轨迹的路径及SHA256。')
    for text in [
        '4.1：main_case_metrics.csv、main_group_summary.csv、main_paired_effects.csv；新增场景为supplement_case_metrics.csv及supplement_group_summary.csv。',
        '4.2：main_local_diagnostics.csv、calibration_case_ellipses.csv、calibration_coverage_summary.csv及校准协议。',
        '4.3：controlled_release_case_metrics.csv、controlled_release_summary.csv、controlled_release_same_event_comparison.csv、literature_metrics_and_conditions.csv。',
        '释放布局比较与共享响应加速消融已从当前实验中移除并可恢复归档；三种方法共用前向响应这一公平比较条件仍保留。']:
        paragraph(document,text,size=10)
    heading(document, '主要文献与比较条件')
    for text in [
        'Vänskä等 2025，Inverse Problems 41 095001。doi:10.1088/1361-6420/adfb29，实验Case2表2。本文1min平均、已知单源；文献3–5s状态更新、分布式源图与地面层源先验。',
        'Hirst等 2020，GRL。doi:10.1029/2019GL086725。原始数据doi:10.5286/edata/738；17次释放合并为不同配置，不能与逐事件中位数直接排名。',
        'Voss等 2024，Atmospheric Environment X。doi:10.1016/j.aeaoa.2024.100260。公开报告3–12m定位、80%源强误差低于30%；保留其统计口径。',
        'Newman等 2025，AOAS。doi:10.1214/25-AOAS2101。浓度RMSE不是定位误差。其2025预印本arXiv:2506.14597报告粒子平均距离，也不同于点估计误差；作为背景参照，不混入同指标排名。']:
        paragraph(document,text,size=9.5)
    destination = packet/'论文实验结果更新汇总.docx'
    document.save(destination)
    print(destination)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--packet',type=Path,default=PACKET)
    parser.add_argument('--date',dest='report_date')
    args = parser.parse_args()
    build(args.packet,args.report_date)
