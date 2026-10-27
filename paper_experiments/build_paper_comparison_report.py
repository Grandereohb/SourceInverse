"""Build the replacement chapter report from complete new-baseline results.

The old report and results are never overwritten. --check-inputs lists actual
missing deliverables without creating a document or inventing placeholder data.
"""

import argparse
from datetime import date
import hashlib
import json
from pathlib import Path

import numpy as np
import pandas as pd
from analyze_paper_newman_batch import chain_summary


ROOT = Path(__file__).resolve().parents[1]/'result/paper_experiments_20261017'
METHODS = dict(ia_vpdi='IA-VPDI', joint_nls_trf='Joint NLS',
               standard_gbest_pso='标准 PSO', matched_joint_adam='联合 Adam')
NETWORKS = dict(ideal5='理想五站', jjj6='实际六站几何',
                surround5_pulse='环绕五站几何', corridor6_double_peak='走廊六站几何')
FIGURES = ('synthetic_localization', 'synthetic_release_cost',
           'localization_reliability', 'public_release_localization')


def check_inputs(synthetic, public, newman, figures, allow_numerical_warnings=False):
    missing, sources = [], []
    protocols = {}
    for label, folder in [('synthetic', synthetic), ('public', public)]:
        protocol = json.loads((folder/'protocol.json').read_text(encoding='utf-8'))
        protocols[label] = protocol
        status = json.loads((folder/'analysis/analysis_status.json').read_text(encoding='utf-8'))
        if not status['final']:
            missing.append(label+' batch incomplete')
        metrics = folder/'analysis/case_metrics.csv'
        if not metrics.exists():
            missing.append(label+' metrics missing')
            continue
        frame = pd.read_csv(metrics)
        if label == 'synthetic':
            expected = {item['case']['case_id'] for item in protocol['cases']}
            actual = set(frame.case_id)
            groups = frame.groupby('case_id')
        else:
            expected = {(item['source'].replace('_', '')+'_'+item['release'], height)
                        for item in protocol['events'] for height in protocol['heights_m']}
            actual = set(zip(frame.event, frame.height_m))
            groups = frame.groupby(['event', 'height_m'])
        if actual != expected:
            missing.append(label+' case identities incomplete')
        expected_methods = set(protocol.get('methods', METHODS))
        if any(set(group.method) != expected_methods or len(group) != len(expected_methods)
               for _, group in groups):
            missing.append(label+' per-case method rows incomplete or duplicated')
        sources.append(metrics)
    replay_status_path = public/'analysis/quadrature_replay_status.json'
    if not replay_status_path.exists():
        missing.append('public higher-quadrature replay missing')
    else:
        replay_status = json.loads(replay_status_path.read_text(encoding='utf-8'))
        if not replay_status['final'] or (replay_status['failed'] and not allow_numerical_warnings):
            missing.append('public higher-quadrature replay incomplete or failed')
        sources.extend((replay_status_path, public/'analysis/quadrature_replay.csv'))
    for name in ('local_diagnostics.csv', 'calibrated_ellipses.csv', 'coverage_summary.csv',
                 'calibration.json', 'risk_coverage.csv', 'sampled_alternatives.csv'):
        path = synthetic/'analysis'/name
        if not path.exists():
            missing.append(name+' missing')
        else:
            sources.append(path)
    alternatives_status = synthetic/'analysis/alternatives_status.json'
    if not alternatives_status.exists() or not json.loads(alternatives_status.read_text(encoding='utf-8'))['final']:
        missing.append('finite candidate-pool diagnostics incomplete')
    if protocols['synthetic'].get('q_capacity_selection') == 'bic':
        ablation = synthetic.parent.parent/'solver_upgrade_full_ia_v2/comparison.csv'
        if not ablation.exists():
            missing.append('fixed-capacity reduced-TRF ablation missing')
        else:
            sources.append(ablation)
        control = synthetic.parent/'capacity_control'
        if not (control/'status.json').exists() or not json.loads((control/'status.json').read_text())['final']:
            missing.append('matched release-capacity control incomplete')
        else:
            design = json.loads((control/'protocol.json').read_text(encoding='utf-8'))
            control_frame = pd.read_csv(control/'comparison.csv')
            expected = {item['case']['case_id'] for item in design['cases']}
            if (set(control_frame.case_id) != expected or len(expected) != 24
                    or any(len(g) != 2 or set(g.method) != {'ia_vpdi_bic', 'joint_nls_bic'}
                           for _, g in control_frame.groupby('case_id'))):
                missing.append('matched release-capacity control rows incomplete or duplicated')
            sources.extend((control/'comparison.csv', control/'protocol.json'))
    newman_records = []
    for event, height in sorted({(item['source'].replace('_', '')+'_'+item['release'], h)
            for item in protocols['public']['events'] for h in protocols['public']['heights_m']}):
        folder = newman/f'{event}_h{int(round(height*10)):02d}'
        path, audit_path = folder/'result.json', folder/'chain_audit.json'
        if not path.exists() or not audit_path.exists():
            missing.append(f'Newman result/audit missing: {folder.name}')
            continue
        record = json.loads(path.read_text(encoding='utf-8'))
        audit = json.loads(audit_path.read_text(encoding='utf-8'))
        if (record['event'] != event or record['height_m'] != height or
                record['iterations'] < 10000 or not record['finite_required_chains'] or
                'posterior_mean_q_kg_h' not in record):
            missing.append('Newman invalid full-window result: '+folder.name)
        if audit['result_sha256'] != hashlib.sha256(path.read_bytes()).hexdigest():
            missing.append('Newman audit does not match result: '+folder.name)
        record['audit'] = audit
        newman_records.append(record)
        sources.extend((path, audit_path))
    for name in FIGURES:
        for suffix in ('.png', '.pdf', '.alignment.json', '.collisions.json'):
            if not (figures/(name+suffix)).exists():
                missing.append('figure/QA missing: '+name+suffix)
        collision = figures/(name+'.collisions.json')
        if collision.exists():
            report = json.loads(collision.read_text(encoding='utf-8'))
            if any(item.get('severity', '').upper() == 'FAIL' for item in report.get('findings', [])):
                missing.append('figure collision failure: '+name)
    manifest = figures/'figure_manifest.json'
    if not manifest.exists() or json.loads(manifest.read_text(encoding='utf-8'))['preview']:
        missing.append('formal figure manifest absent or marked preview')
    return dict(ready=not missing, missing=missing,
                source_files=[str(path) for path in sources], newman_records=newman_records)


def interval(series, digits=1):
    values = np.asarray(series.dropna(), dtype=float)
    if not len(values):
        return '无实测参考'
    low, center, high = np.quantile(values, [.25, .5, .75])
    return f'{center:.{digits}f} [{low:.{digits}f}, {high:.{digits}f}]'


def build(synthetic, public, newman, figures, output, allow_numerical_warnings=False):
    evidence = check_inputs(synthetic, public, newman, figures, allow_numerical_warnings)
    if not evidence['ready']:
        raise ValueError('Report inputs incomplete: '+json.dumps(evidence['missing'], ensure_ascii=False))
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    from build_paper_results_summary import font, paragraph, heading, table

    frame = pd.read_csv(synthetic/'analysis/case_metrics.csv')
    beam = pd.read_csv(public/'analysis/case_metrics.csv')
    public_methods = {k: v for k, v in METHODS.items() if k in set(beam.method)}
    updated = json.loads((synthetic/'protocol.json').read_text(encoding='utf-8')).get('q_capacity_selection') == 'bic'
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.7), Cm(1.8)
    section.left_margin, section.right_margin = Cm(1.8), Cm(1.8)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
        style = document.styles[name]
        style.font.name, style.font.color.rgb = 'Arial', RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        if name == 'Normal':
            style.font.size = Pt(10.5)
        properties = style._element.find(qn('w:pPr'))
        if properties is not None:
            for border in list(properties.findall(qn('w:pBdr'))):
                properties.remove(border)
    footer = section.footer.paragraphs[0]
    footer.alignment = 1
    font(footer.add_run('第四章新基线实验报告  '), 9)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    title = document.add_paragraph(style='Title')
    font(title.add_run('第四章新基线实验报告'), 21, True)
    paragraph(document, date.today().strftime('%Y年%m月%d日')+'  供论文写作核对', size=10)
    paragraph(document, '本报告采用既有实验输入，更新IA-VPDI与Joint NLS、标准PSO的共同正演比较，'
        '保留联合Adam消融，并使用Newman等2025年作者Gaussian plume及M-MALA-within-Gibbs进行公开释放比较。'
        '下面按实验目的、测试内容、对比方法和数据结果组织。'+
        ('新版IA使用尺度化约化信赖域及常量、三帽、五帽释放容量选择；完整保留旧对照和全部案例。'
         if updated else '未新增算法优化实验。'))
    paragraph(document, '章节映射：本报告4.2、4.3、4.4对应原实验计划中的4.1合成评价、4.2可靠性、'
        '4.3公开释放；4.1统一列出实验设置。历史结果保留，不与新边界下的数据混合排名。', size=10)
    heading(document, '4.1 实验设置与评价方法')
    table(document, ['方法', '待估变量', '求解方式'], [
        ['IA-VPDI', '位置 扩散与非负Q系数',
         'Q消元 约化TRF及BIC容量选择' if updated else '内层非负最小二乘 外层Adam'],
        ['Joint NLS', '位置 扩散与非负Q系数', 'SciPy TRF联合约束最小二乘'],
        ['标准 PSO', '位置 扩散与非负Q系数', 'PySwarms固定系数GlobalBestPSO'],
        ['联合 Adam', '位置 扩散与非负Q系数', '共同正演的联合优化消融'],
        ['Newman 2025', '位置 排放 扩散及背景等', '作者Gaussian plume及M-MALA-within-Gibbs']], [3.0, 5.4, 9.0])
    if updated:
        paragraph(document, '216例使用相同观测、共同五帽函数输运空间、有限参数域和两个原共同粗初值。'
            'IA在常量、三帽、五帽嵌套子空间分别用约化TRF求解，以Gaussian工作模型BIC选择容量，'
            '每模型每起点最多250次函数评价；原固定五参数NLS和联合Adam保留。'
            '另设每组几何按冻结协议顺序取6例的24例同容量NLS对照，区分模型选择与求解器收益。'
            'BIC使用全部观测，不按真实释放类型选容量；本批次为开发性完整复核，不另称全新独立测试集。')
        paragraph(document, '标准PSO采用PySwarms固定系数GlobalBestPSO，每种子32粒子256代、3个预设种子。'
            '75%粒子使用共同初值及局部扰动，其余全域初始化；按物理尺度限速、零初速度。'
            '粒子群前向使用单精度加速，最终以双精度原正演回放，目标值相对差须不超过0.01%。'
            '不使用NNLS消元或NLS局部精修。IA多容量增加预算，不能把记录耗时视为隔离等预算基准。')
    else:
        paragraph(document, '共同正演比较：216例使用相同观测、5个hat源强基、有限参数域和两个原共同粗初值。'
        'IA及Adam各起点1000步；NLS各起点最多500次函数评估，并计实际Jacobian调用；'
        'PSO每种子32粒子64代，共3个预设种子，保留全部重复。不同更新次数不等价FLOPs；'
            '报告时间为复用共同粗初值后的条件求解时间，不包括重新执行全域粗搜索。')
    paragraph(document, 'PSO先在每个案例内对三种子的指标取中位数，再汇总案例，不取最好种子。'
        '所有位置边界来自既有搜索场地，源强上界由观测初始化确定，不使用真源位置或真源强。'
        '定位误差为估计与已知源的平面欧氏距离，源强重建单独评价。表中中心与区间均为中位数和四分位区间。')

    heading(document, '4.2 合成场景定位与源强重建')
    paragraph(document, '目的：在正演和输入一致时比较不同求解策略，并考察既有站网、风场与释放过程中的表现。'
        '192例名义条件覆盖四个场景，每场景48例；另保留12例风场失配和12例源强形状失配。'
        '实际六站仅沿用真实站网几何，浓度、源点和源强真值仍由合成输入给定。')
    if updated:
        paragraph(document, '合成真值采用独立实现对连续释放的各向异性平流扩散响应做高精度积分，'
            '包含时变风、非负源强和衰减，并完成时间与释放求积加密核查。'
            '六个小时平均浓度按峰值归一后加入标准差0.01的高斯噪声；源点与源强真值由生成过程保存，'
            '不从高值站推定。反演使用共同的输运响应与观测平均算子；名义条件检验匹配物理假设下的求解能力，'
            '风场与释放形状失配另列。合成Q为模型单位，只比较相对重建误差，不解释为实测kg/h。', size=9)
        paragraph(document, '理想五站与实际六站各为两类风场×三类释放过程×8次随机重复；'
            '环绕五站固定为稳态风下单峰释放48例，走廊六站固定为转向风下双峰释放48例。'
            '后两组是专门场景，不把四组共同描述为完整因子设计。', size=9)
    nominal = frame[frame.cohort == 'nominal']
    for network, label in NETWORKS.items():
        subset = nominal[nominal.network == network]
        paragraph(document, label+'  n='+str(subset.case_id.nunique()), bold=True)
        table(document, ['方法', '定位误差 m', 'Q相对L2误差 %', '求解时间 s'], [
            [label_method, interval(group.localization_error_m), interval(100*group.q_relative_l2),
             interval(group.seconds)] for method, label_method in METHODS.items()
            for group in [subset[subset.method == method]]], [3.1, 4.8, 4.8, 4.7])
    paragraph(document, '全部216例的定位误差中位数分别为：'+ '；'.join(
        f'{label} {frame.loc[frame.method == method, "localization_error_m"].median():.1f} m'
        for method, label in METHODS.items())+'。这组总体统计仅用于描述固定样本；'
        '场景同时改变几何、风场或释放过程，不作为单一因素的因果效应。', bold=True)
    if updated:
        table(document, ['方法', '定位中位 m', '定位均值 m', 'Q中位误差 %', 'Q均值误差 %'], [
            [label, f'{group.localization_error_m.median():.2f}',
             f'{group.localization_error_m.mean():.2f}', f'{100*group.q_relative_l2.median():.2f}',
             f'{100*group.q_relative_l2.mean():.2f}'] for method, label in METHODS.items()
            for group in [frame[frame.method == method]]], [3.0, 3.5, 3.5, 3.7, 3.7])
    document.add_picture(str(figures/'synthetic_localization.png'), width=Cm(17.1))
    paragraph(document, '图1 各场景定位误差的经验分布。包括全部216例，失配条件单独显示；PSO使用案例内三种子中位数。', size=9)
    mismatch = frame[frame.cohort != 'nominal']
    table(document, ['失配条件', '方法', '定位误差 m', 'Q相对L2误差 %'], [
        [('风场失配' if cohort == 'wind_mismatch' else '释放形状失配'), METHODS[method],
         interval(group.localization_error_m), interval(100*group.q_relative_l2)]
        for (cohort, method), group in mismatch.groupby(['cohort', 'method'])], [3.4, 3.1, 5.5, 5.4])
    document.add_picture(str(figures/'synthetic_release_cost.png'), width=Cm(17.1))
    paragraph(document, '图2 全部案例的源强误差及相对IA的配对求解时间。时间比显示中位数与四分位区间，'
        'PSO时间为单次种子的中位数而非三次总和。', size=9)
    if updated:
        paragraph(document, '释放容量与求解策略消融', bold=True)
        ablation = pd.read_csv(synthetic.parent.parent/'solver_upgrade_full_ia_v2/comparison.csv')
        variants = [('旧IA Adam 固定五参数', ablation[ablation.method == 'ia_vpdi']),
                    ('IA 约化TRF 固定五参数', ablation[ablation.method == 'ia_vpdi_reduced_trf']),
                    ('IA 约化TRF 容量选择', frame[frame.method == 'ia_vpdi'])]
        table(document, ['完整216例 IA版本', '定位误差 m', 'Q相对L2误差 %', '条件求解 s'], [
            [label, interval(group.localization_error_m), interval(100*group.q_relative_l2),
             interval(group.seconds)] for label, group in variants], [5.0, 4.2, 4.4, 3.8])
        if 'q_basis_count' in frame:
            counts = frame.loc[frame.method == 'ia_vpdi', 'q_basis_count'].value_counts()
            paragraph(document, '新版IA选择常量、三帽、五帽释放模型分别为'+
                '、'.join(str(int(counts.get(k, 0))) for k in (1, 3, 5))+'例。'
                '上表复用既有完整结果，不新增优化；各版本预算及执行时段不同，耗时作为条件记录。', size=9)
            selected = frame[(frame.method == 'ia_vpdi') & (frame.cohort == 'nominal')]
            constant = selected[selected.release == 'constant']
            nonconstant = selected[selected.release != 'constant']
            counts_constant = constant.q_basis_count.value_counts()
            paragraph(document, f'名义条件下，{len(constant)}例真实恒定释放分别有'+
                '、'.join(str(int(counts_constant.get(k, 0))) for k in (1, 3, 5))+
                '例选择常量、三帽、五帽模型；'+
                f'{int((nonconstant.q_basis_count == 5).sum())}/{len(nonconstant)}例单峰或双峰释放保留五参数。'
                '真实释放类型只用于事后解释，未输入容量选择规则。', size=9)
        control = pd.read_csv(synthetic.parent/'capacity_control/comparison.csv')
        table(document, ['同容量方法 n=24', '定位误差 m', 'Q相对L2误差 %', '条件求解 s'], [
            [('IA 容量选择' if method == 'ia_vpdi_bic' else '联合NLS 容量选择'),
             interval(group.localization_error_m), interval(100*group.q_relative_l2), interval(group.seconds)]
            for method, group in control.groupby('method')], [5.0, 4.2, 4.4, 3.8])
        paragraph(document, '该消融给两种求解器相同容量族、BIC、观测、边界、几何初值及评价上限；'
            'NLS只在共同初值处估计Q，此后联合优化而不消元。固定五参数对照与容量选择对照分开解释，'
            '不把模型复杂度选择带来的收益全部归因为变量投影。24例按冻结元数据顺序预定，每组几何6例；'
            '前两组覆盖六种风场/释放组合，专门单峰和双峰场景各取六次重复，不按胜负挑选。', size=9)
        paired_objectives = control.pivot(index='case_id', columns='method', values='observed_mse')
        near = int(((paired_objectives.ia_vpdi_bic/paired_objectives.joint_nls_bic-1).abs() <= .01).sum())
        ia_time = control.loc[control.method == 'ia_vpdi_bic', 'seconds'].median()
        nls_time = control.loc[control.method == 'joint_nls_bic', 'seconds'].median()
        paragraph(document, f'同容量对照中，{near}/{len(paired_objectives)}例目标差不超过1%；'
            f'IA与NLS的条件求解时间中位数分别为{ia_time:.2f}s和{nls_time:.2f}s。'
            '精度相近时，约化求解体现计算效率；容量适配的重建收益由上表的版本消融解释。', bold=True, size=9)

    heading(document, '4.3 定位可靠性与区域校准')
    paragraph(document, '目的：检验消除源强和扩散局部补偿后的区域是否能反映定位不确定性，而非只比较浓度拟合。'
        '在原两类站网96个名义案例上，比较固定Q与D、只消除Q、同时消除Q与D三种局部信息；'
        '继续采用48例校准及48例验证的既定案例划分，不从反演中删除站点或时段。')
    if updated:
        paragraph(document, '本节已按新版IA逐案例实际选中的释放容量和源点重算。局部区域及有限候选池分析均条件于所选容量，'
            '不包含离散模型选择不确定性。不可定义椭圆的案例保留在分母中、不给予覆盖，校准时记为正无穷分数；'
            '既定48/48划分属于回顾性校准，而非新增独立外部验证。', size=9)
    local = pd.read_csv(synthetic/'analysis/local_diagnostics.csv')
    local = local[(local.cohort == 'nominal') & local.network.isin(['ideal5', 'jjj6'])]
    table(document, ['局部信息', '包含真源 例数', '可定义椭圆 例数', '面积中位数 km2'], [
        [label, f'{int(group.contains_truth.eq(True).sum())}/96', f'{int(group.available.sum())}/96',
         interval(group.area_m2/1e6, 3)]
        for key, label in [('naive_fixed_Q_D', '固定 Q与D'), ('eliminate_Q', '消除 Q'),
                           ('eliminate_Q_D', '消除 Q与D')]
        for group in [local[local.diagnostic == key]]], [4.2, 4.0, 4.4, 4.8])
    coverage = pd.read_csv(synthetic/'analysis/coverage_summary.csv')
    table(document, ['验证分组', '原包含率 %', '校准包含率 %', '校准面积 km2'], [
        [{'all': '全部48例', 'ideal5': '理想五站24例', 'jjj6': '实际六站24例'}[row.network],
         f'{100*row.raw_coverage:.1f}', f'{100*row.calibrated_coverage:.1f}',
         f'{row.calibrated_area_median_m2/1e6:.3f}']
        for row in coverage[coverage.partition == 'validation'].itertuples()], [4.4, 4.0, 4.0, 5.0])
    document.add_picture(str(figures/'localization_reliability.png'), width=Cm(17.1))
    paragraph(document, '图3 上排为同一48例验证集的经验包含率与区域面积；下排为两类站网各48例按估计长半轴'
        '排序后的保留比例与平均定位误差。曲线为描述性结果，不用于删除反演观测或修改源点。'
        '包含率应结合面积解释，不能单独当作点定位精度提高。', size=9)
    paragraph(document, '近等价候选位置', bold=True)
    paragraph(document, '在原96例既有有限候选池上，以新共同参数边界重新求条件最优Q，不新增位置或外层优化。'
        '以已评估最低目标为参考，保留1%、5%、10%阈值内位置。若池中发现更低目标，单独标记当前解的差距，'
        '不自动替换源点。有限网格未发现备选位置不等于真实位置唯一，也不构成连续置信域。')
    alternatives = pd.read_csv(synthetic/'analysis/sampled_alternatives.csv')
    table(document, ['目标容差', '多位置例数', '最远距离中位 m', '最远距离最大 m', '当前解在阈值外'], [
        [f'{100*tolerance:.0f}%', f'{int((group.sampled_locations > 1).sum())}/{group.case_id.nunique()}',
         f'{group.farthest_alternative_m.median():.1f}', f'{group.farthest_alternative_m.max():.1f}',
         str(int((~group.selected_within_tolerance).sum()))]
        for tolerance, group in alternatives.groupby('relative_score_tolerance')], [1.8, 3.2, 4.0, 4.0, 4.4])

    heading(document, '4.4 公开受控释放数据验证')
    paragraph(document, '目的：利用Chilbolton甲烷受控释放检验真实观测下的源点与排放反演。'
        '包含8个释放窗口、2个物理源位置及7条光路；两种记录源高0.3m和1.6m分别报告，不择优。'
        '本轮IA、NLS和Adam共用垂直扩散、地面反射、光路平均及ppm转换适配器，均采用常Q容量；'
        '旧五基函数结果仅保留为容量不同的补充材料。')
    if updated:
        paragraph(document, '本节IA更新为约化TRF，每起点最多250次评价；已知恒定受控释放固定采用常Q，'
            '不强行进行时变容量选择。原NLS和Adam完整结果只读复用；本节另与作者贝叶斯方法比较。'
            '新版PSO已在216例合成实验中评价，为控制计算量，本轮不再重复公开PSO；旧版完整结果单独保留，'
            '不混入新版方法排名。两种源高仍分别报告，未选择效果较好的设定。')
    else:
        paragraph(document, '本节预算：IA及联合Adam各起点500步，NLS各起点最多250次函数评估，'
        'PSO每种子32粒子32代，均使用相同两个共同粗初值。公开比较统一使用CUDA结果，'
            '原CPU参考只用于数值一致性核查，不混合汇总运行时间。')
    replay = pd.read_csv(public/'analysis/quadrature_replay.csv')
    paragraph(document, f'数值核查：对全部{len(replay)}项拟合输出，将时间及释放求积由16阶提高到32阶，'
        '在原拟合参数处回放而不再优化。相对预测差1%作为既有阈值，'
        f'通过{int(replay.gate_passed.sum())}/{len(replay)}项，最大相对差'
        f'{100*replay.relative_prediction_change.max():.3f}%。该核查只证明本次有限求积的一致性，不证明源位置正确。')
    failed_replay = replay[~replay.gate_passed]
    if len(failed_replay):
        paragraph(document, '数值精度例外：以下输出未通过既定1%阈值，未放宽阈值、删除结果或重新优化。'
            '这些条目仍纳入完整样本；本报告不将其表述为全批次数值核查通过。'
            '同一事件的微小方法差别不宜据此判定优劣，后续如需精修应单独版本化。', bold=True)
        table(document, ['事件', '高 m', '方法', '相对预测差 %'], [
            [row.event, row.height_m, METHODS[row.method], f'{100*row.relative_prediction_change:.3f}']
            for row in failed_replay.itertuples()], [5.0, 2.0, 5.4, 5.0])
    for height in (.3, 1.6):
        paragraph(document, f'共同正演比较 源高{height:g} m', bold=True)
        table(document, ['方法', '8窗口定位误差 m', '实测Q相对误差 %', '浓度RMSE ppm'], [
            [label, interval(group.localization_error_m), interval(100*group.q_relative_error),
             interval(group.observed_rmse_ppm, 3)] for method, label in public_methods.items()
            for group in [beam[(beam.height_m == height) & (beam.method == method)]]], [3.1, 5.0, 5.0, 4.3])
    document.add_picture(str(figures/'public_release_localization.png'), width=Cm(17.1))
    paragraph(document, '图4 共同正演下逐事件定位误差。'+
        ('' if updated else 'PSO浅色点为全部随机重复，主符号为种子中位数；')+
        '源高敏感性和同一物理源的多个窗口不计作独立源位置。', size=9)
    paragraph(document, 'Newman作者方法', bold=True)
    paragraph(document, '调用sourceinversion 0.1.5作者Gaussian plume、原后验和M-MALA-within-Gibbs，'
        '不以其他采样器替代。每项10000步，丢弃前5000步；共同预窗口残差加1.92ppm参考背景，由作者方法估计背景；'
        '两源作者notebook相同的位置先验保留，后验与LHS似然函数经AST核对一致。'
        '初始化及单窗口适配与原文合并窗口流程有差异，逐次记录。'
        '因此这属于方法级比较，不是同正演优化器消融或原文表格的严格复现。')
    table(document, ['窗口', '高 m', '位置误差 m', '估计Q kg/h', '位置最小ESS'], [
        [record['event'], record['height_m'], f'{record["localization_error_m"]:.2f}',
         f'{record["posterior_mean_q_kg_h"]:.3f}',
         (f'{chain_summary(record, record["audit"])["position_min_ess"]:.1f}'
          if chain_summary(record, record['audit'])['position_min_ess'] is not None else '未定义')]
        for record in evidence['newman_records']], [4.5, 1.5, 3.6, 3.8, 4.0])
    for height in (.3, 1.6):
        records = [item for item in evidence['newman_records'] if item['height_m'] == height]
        paragraph(document, f'Newman固定预算摘要 源高{height:g}m：8窗口定位误差中位数'
            f'{np.median([item["localization_error_m"] for item in records]):.2f}m，'
            f'参数均值处浓度回放RMSE中位数'
            f'{np.median([item["audit"]["parameter_mean_forward_rmse_ppm"] for item in records]):.3f}ppm。')
    paragraph(document, 'Newman表格采用固定预算链的后验均值摘要，同时列出ESS；未通过链混合检查的条目'
        '不作为已收敛方法的精度结论或校准区间证据。浓度回放是在物理参数均值处计算，'
        '并非后验预测均值。独立称重缺失的事件不计算实测源强误差。')
    heading(document, '数据复核与写作使用')
    paragraph(document, '全部新数据保留逐案例、逐随机种子及运行预算记录；旧NNLS-Powell和已发表结果'
        '只作补充参照。论文定量比较应使用此版完整样本，不按算法胜负删样。定位误差、源强重建、'
        '浓度拟合和可靠性区域分别报告；计算成本结论限定于实际记录的预算与硬件条件。')
    output = Path(output)
    if output.exists():
        raise FileExistsError('Choose a new report output; do not overwrite the prior report')
    output.parent.mkdir(parents=True, exist_ok=True)
    # Keep table labels and initial rows together without forcing long tables onto one page.
    for item in document.paragraphs:
        following = item._p.getnext()
        if (following is not None and following.tag == qn('w:tbl')) or item._p.findall('.//' + qn('w:drawing')):
            item.paragraph_format.keep_with_next = True
    for grid in document.tables:
        for row in grid.rows:
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        initial_rows = grid.rows if len(grid.rows) <= 9 else grid.rows[:2]
        for row in initial_rows:
            for cell in row.cells:
                for item in cell.paragraphs:
                    item.paragraph_format.keep_with_next = row is not initial_rows[-1]
    document.save(output)
    (output.with_suffix('.sources.json')).write_text(json.dumps(dict(
        generated=date.today().isoformat(), experiments_reoptimized=bool(updated),
        source_sha256={str(Path(path).resolve().relative_to(ROOT.resolve())): hashlib.sha256(Path(path).read_bytes()).hexdigest()
                       for path in evidence['source_files']}), ensure_ascii=False, indent=2), encoding='utf-8')
    return output


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--synthetic', type=Path, default=ROOT/'new_comparison_v4')
    parser.add_argument('--public', type=Path, default=ROOT/'beam_comparison_cuda_v1')
    parser.add_argument('--newman', type=Path, default=ROOT/'newman_lhs_full_v1')
    parser.add_argument('--figures', type=Path, default=ROOT/'new_comparison_v4/figures')
    parser.add_argument('--output', type=Path, default=ROOT/f'第四章新基线实验报告_{date.today():%Y%m%d}.docx')
    parser.add_argument('--check-inputs', action='store_true')
    parser.add_argument('--allow-numerical-warnings', action='store_true',
                        help='Report complete numerical checks with explicit failed-threshold rows; does not certify them')
    args = parser.parse_args()
    if args.check_inputs:
        report = check_inputs(args.synthetic, args.public, args.newman, args.figures, args.allow_numerical_warnings)
        print(json.dumps({key: value for key, value in report.items() if key != 'newman_records'},
                         ensure_ascii=False, indent=2))
    else:
        print(build(args.synthetic, args.public, args.newman, args.figures, args.output, args.allow_numerical_warnings))
