"""Four-page, plain-language experiment report from complete audited evidence."""

import argparse
import hashlib
import json
from pathlib import Path

import numpy as np
import pandas as pd

from run_focused_chapter import ROOT, STUDY, UPDATED, OUT, METHODS, NETWORKS, read, save


def build(output, metec=False):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from build_paper_results_summary import paragraph, heading, table, font
    if output.exists():
        raise FileExistsError('Choose a new report output; preserve prior deliveries')
    if not read(OUT/'analysis/summary.json')['final']:
        raise ValueError('Complete analysis required')
    for path, digest in read(OUT/'analysis/sources.json').items():
        if hashlib.sha256((ROOT/path).read_bytes()).hexdigest() != digest:
            raise ValueError('Source changed after analysis: '+path)
    for name in ['case_eval_007', 'case_eval_009', 'fit_and_location', 'public_locations', 'public_case']:
        qa = read(OUT/'figures'/(name+'.collisions.json'))
        if any(item.get('severity', '').upper() == 'FAIL' for item in qa.get('findings', [])):
            raise ValueError('Figure collision failure: '+name)
    frame = pd.read_csv(OUT/'analysis/case_flags.csv')
    nominal = frame[frame.cohort == 'nominal']
    summary = pd.read_csv(OUT/'analysis/nominal_summary.csv').set_index('method')
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.5), Cm(1.6)
    section.left_margin, section.right_margin = Cm(1.8), Cm(1.8)
    for name in ('Normal', 'Title', 'Heading 1'):
        style = document.styles[name]
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.name = 'Arial'
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        properties = style._element.find(qn('w:pPr'))
        if properties is not None:
            for border in list(properties.findall(qn('w:pBdr'))):
                properties.remove(border)
    footer = section.footer.paragraphs[0]
    footer.alignment = 1
    font(footer.add_run('污染源定位实验简报  '), 9)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    def p(text, bold=False):
        item = paragraph(document, text, bold, 10)
        item.paragraph_format.line_spacing = 1.04
        item.paragraph_format.space_after = Pt(4)
        return item
    def image(name, caption, width=17.):
        document.add_picture(str(OUT/'figures'/(name+'.png')), width=Cm(width))
        item = paragraph(document, caption, size=9)
        item.paragraph_format.line_spacing = 1.04
        item.paragraph_format.space_after = Pt(4)
    def grid(headers, rows):
        table(document, headers, rows, [17.4/len(headers)]*len(headers))
        blank = document.paragraphs[-1]._p
        blank.getparent().remove(blank)
    def med(series, scale=1, digits=1):
        return f'{scale*series.median():.{digits}f}'
    def center_pair(series, scale=1, digits=1):
        return f'{scale*series.median():.{digits}f} / {scale*series.mean():.{digits}f}'
    def new_page_heading(text):
        heading(document, text)
        document.paragraphs[-1].paragraph_format.page_break_before = True
    title = document.add_paragraph(style='Title')
    font(title.add_run('污染源定位实验简报'), 20, True)
    paragraph(document, '2026年10月27日  按精简后的实验提纲整理', size=9)
    heading(document, '4.1 实验设置与评价指标')
    p('研究问题：IA能否更准确地找到源点，并减少“浓度拟合很好、位置却不对”的情况。'
      '保留全部预定案例；不新增初值敏感性或机制组合实验。')
    p('比较IA-VPDI、联合最小二乘（Joint NLS）、标准粒子群（PSO），联合Adam作为补充。'
      '合成真值由独立输运计算产生并保存；原理想/实际站网各有两类风场、三类释放各8次，'
      '环绕/走廊组分别固定为单峰/双峰释放各48次，并非四组完整因子设计。'
      '实际六站组只使用真实站网几何，不是真实泄漏观测。')
    p('正常条件共192例；输入不准确的24例另外保留。各方法使用相同观测、物理模型和搜索边界。'
      '完整IA自动选择用1、3或5个参数描述排放，其他主对照固定5个。'
      'PSO保留3个种子，先在案例内取指标中位数；不同预算不称等计算量比较。')
    heading(document, '4.2 合成场景中的定位与源强重建')
    p('目的：比较源点和排放过程的重建。汇总表中“中位数 / 平均数”同时报告：前者反映典型案例，'
      '后者体现少数严重偏差的影响；单位见表头，均越小越好。')
    rows = []
    for network, label in NETWORKS.items():
        rows.append([label+'（48例）', *[center_pair(nominal[(nominal.network == network) &
            (nominal.method == method)].localization_error_m) for method in METHODS]])
    rows += [['192例定位误差', *[center_pair(nominal[nominal.method == m].localization_error_m, digits=2) for m in METHODS]],
             ['192例源强误差%', *[center_pair(nominal[nominal.method == m].q_relative_l2, 100, 2) for m in METHODS]]]
    grid(['指标或场景', *METHODS.values()], rows)
    p('IA在理想五站和实际六站组的定位中位误差较低；环绕组NLS略低，走廊组基本一致。'
      '192例中位定位误差IA与NLS同为21.53米，平均数分别为72.11与104.04米；差异主要体现在较大的偏差，见4.3。'
      '不能由此表述为所有场景全面领先。', True)
    grid(['补充场景', *METHODS.values()], [[label+'（12例）', *[center_pair(frame[
        (frame.cohort == cohort) & (frame.method == m)].localization_error_m) for m in METHODS]]
        for cohort,label in [('wind_mismatch', '风场有偏差'), ('q_mismatch', '排放形状不匹配')]])
    p('风场偏差为风速±10%、风向±5°；形状不匹配指真实排放用不规则6节点生成，反演只用固定节点。'
      '两组均未显示IA明显领先。全部216例仍保留，IA/NLS总体定位中位数为25.66/28.14米；'
      '正文不把不同条件混在一起解释为同一种效果。')

    new_page_heading('4.2 合成代表案例')
    case = 'eval_007'
    record = read(UPDATED/'synthetic'/case/'result.json')
    image('case_'+case, '图1 实际六站几何下的合成案例eval_007。左：黑星为真源、灰点为站点；'
        '中：真实和重建排放；右：观测峰值最高站的浓度。不是实际排放事件。')
    grid(['方法', '定位误差 米', '源强误差%'], [[METHODS[m], f'{record["methods"][m]["localization_error_m"]:.1f}',
        f'{100*record["methods"][m]["q_relative_l2"]:.1f}'] for m in ['ia_vpdi', 'joint_nls_trf']])
    p('此案例按冻结顺序预先选取，未按算法胜负挑选。线或点重合表示结果接近，不代表缺少输出。')
    heading(document, '4.3 完整方法的错误定位对比')
    p('检验办法：把浓度拟合与真实定位误差放在一起看。“拟合达标”暂定为浓度误差不超过已知噪声标准差的2倍'
      '（归一化误差0.02）；“位置偏差较大”暂定为超过100米。这是预先固定的分析口径，不是仪器或业务标准。')
    grid(['方法', '拟合达标例数', '达标但偏位例数', '占全部192例%', '定位误差90分位 米'], [
        [label, f'{int(summary.loc[m,"good_fit"])}/192', str(int(summary.loc[m,'good_fit_wrong'])),
         f'{100*summary.loc[m,"joint_rate"]:.1f}', f'{summary.loc[m,"p90_location_m"]:.1f}']
        for m,label in METHODS.items()])
    p('IA比NLS少5例“拟合达标但偏位”（34对39例，差2.6个百分点）。'
      '这里的完整IA还包含排放参数数量选择，不能把差别全归因于源强消元。PSO达标案例内的偏位比例为'
      f'{100*summary.loc["standard_gbest_pso","conditional_rate"]:.1f}%，与占全部案例的比例分开。')

    new_page_heading('4.3 浓度拟合与定位误差的关系')
    image('fit_and_location', '图2 每个点是一个正常条件案例；全部192例均展示。横虚线为100米，竖虚线为拟合误差0.02。'
        '左上区域代表“浓度拟合达标，但位置偏差较大”。PSO的两个坐标分别是案例内3种子的指标中位数，非单次解。', width=14.8)
    sensitivity = pd.read_csv(OUT/'analysis/threshold_sensitivity.csv')
    grid(['定位门槛 米', *METHODS.values()], [[int(t), *[str(int(sensitivity[(sensitivity.method == m) &
        (sensitivity.location_threshold_m == t)].good_fit_wrong.iloc[0])) for m in METHODS]] for t in [50,100,200]])
    p('表内是拟合达标但偏位的例数，分母均为192。50米门槛下IA与NLS只差1例；100/200米门槛下分别少5/4例。'
      '方向一致但幅度随门槛变化，不能仅凭100米结果声称显著优势。')
    case = read(OUT/'analysis/summary.json')['representative_ids'][-1]
    record = read(UPDATED/'synthetic'/case/'result.json')
    image('case_'+case, '图3 诊断案例'+case+'：按编号选取NLS首次出现“拟合达标但超过100米”的正常案例，'
        '不要求IA胜出。两种方法在本例都存在位置偏差，说明拟合好不等于位置准确。', width=16.)
    p('本例IA/NLS定位误差分别为'+ '/'.join(f'{record["methods"][m]["localization_error_m"]:.1f}' for m in
      ['ia_vpdi','joint_nls_trf'])+'米。它是现象展示，不单独证明偏差由求解器耦合造成；'
      '也不能把全部偏差都归因于优化失败。')

    if metec:
        from build_metec_report import add_chapter
        add_chapter(document,page_break_before=True)
        output.parent.mkdir(parents=True,exist_ok=True)
        document.save(output)
        from build_metec_report import OUT as metec_out
        save(output.with_suffix('.sources.json'),dict(previous=read(OUT/'analysis/sources.json'),
            metec=read(metec_out/'analysis/sources.json'),
            editorial_change='Removed24-case matched-Q main-text control; archived data unchanged'))
        print(output)
        return
    new_page_heading('4.4 公开受控释放数据中的定位验证')
    p('数据：Chilbolton甲烷受控释放，8个窗口、2个真实源位置、7条光路。0.3米和1.6米源高分别报告，'
      '不是16个独立源点。本轮复用完整数据，不新增释放事件或改动模型。')
    public = pd.read_csv(UPDATED/'public/analysis/case_metrics.csv')
    grid(['源高 米', '方法', '定位误差 米', '实测源强误差%', '浓度误差 ppm'], [[height,METHODS[m],
        center_pair(g.localization_error_m, digits=2), center_pair(g.q_relative_error.dropna(),100,1), center_pair(g.observed_rmse_ppm, digits=3)]
        for height in [.3,1.6] for m in ['ia_vpdi','joint_nls_trf','matched_joint_adam']
        for g in [public[(public.height_m == height) & (public.method == m)]]])
    p('表中为中位数 / 平均数，定位与浓度指标覆盖8窗口；源强指标仅在有独立称重记录的6窗口计算。'
      '这些方法使用相同物理模型及恒定排放假设，结果接近。当前数据尚未显示IA相对NLS的明确定位优势。')
    image('public_case', '图4 真实事件source1_rel03，固定取协议首个事件及0.3米源高，不按结果挑选。'
        '左：黑星是真源，灰线为光路；右：观测峰值最高光路的浓度。各方法估计点和曲线接近重合。'
        '全部8窗口的逐事件比较另存图件public_locations。')
    newman = pd.read_csv(STUDY/'newman_lhs_full_v1/analysis/case_metrics.csv')
    grid(['作者方法独立对照', '定位中位数 / 平均数 米'], [[f'Newman 源高{h:g}米',
        center_pair(g.localization_error_m,digits=2)]
        for h in [.3,1.6] for g in [newman[newman.height_m == h]]])
    p('Newman使用作者羽流模型和采样程序，与共同模型求解器比较分开解释；部分链停滞或有效样本不足，'
      '不据此声称优于已充分收敛的作者方法。数值回放43/48项通过原1%门槛，其余5项最大1.039%，'
      '记录仍保留，源2事件11的细小方法差异不用于判断优劣。')
    p('结论：完整IA在部分合成站网中减少较大的定位误差；相同排放表达及公开恒定释放对照尚未显示'
      '消元本身带来定位精度优势。下一步应针对剩余错误定位的来源改进模型，而非继续压低已经接近的拟合误差。', True)
    for item in document.paragraphs:
        if item._p.findall('.//'+qn('w:drawing')):
            item.paragraph_format.keep_with_next = True
    for grid_item in document.tables:
        for row in grid_item.rows[:-1]:
            for cell in row.cells:
                for item in cell.paragraphs:
                    item.paragraph_format.keep_with_next = True
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    save(output.with_suffix('.sources.json'), read(OUT/'analysis/sources.json'))
    print(output)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, default=OUT/'污染源定位实验简报_20261027.docx')
    parser.add_argument('--metec',action='store_true',help='Use completed full METEC cohort for Chapter4.4')
    args=parser.parse_args()
    build(args.output,args.metec)
