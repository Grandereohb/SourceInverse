from __future__ import annotations

import argparse
import re
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


SCRIPT_DIR = Path(__file__).resolve().parent
REPO_ROOT = SCRIPT_DIR.parent

# ============================= Manual Input =============================
# Edit these values before running this script directly.
INPUT_FILE_PATH = (
    REPO_ROOT / "data" / "jjj" / "2026年小时数据" / "8月小时数据_标准单位_汇总.xlsx"
)
POLLUTANTS: list[str] = ["环己烷", "丙烷", "苯", "甲苯", "甲醇"]
OUTPUT_FILE_PATH = (
    REPO_ROOT / "data" / "jjj" / "2026年小时数据" / "8月小时数据_指定污染物.xlsx"
)
ALLOW_MISSING_POLLUTANTS = False
# ========================================================================

TIME_COLUMN = "时间"
SUPPORTED_SUFFIXES = {".xlsx", ".xlsm"}
UNIT_SUFFIX_PATTERN = re.compile(
    r"\s*[\(（]\s*"
    r"(?:ug|μg|µg|mg|g|ng|ppm|ppb|mol|℃|%|hpa|m/s|°)"
    r"(?:\s*/\s*[^）\)]*)?"
    r"\s*[\)）]\s*$",
    flags=re.IGNORECASE,
)


@dataclass(frozen=True)
class SheetFilterSummary:
    sheet_name: str
    original_columns: int
    retained_columns: int


def resolve_path(path: str | Path) -> Path:
    resolved = Path(path).expanduser()
    if not resolved.is_absolute():
        resolved = REPO_ROOT / resolved
    return resolved.resolve()


def normalize_name(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower()
    text = text.replace("（", "(").replace("）", ")")
    return re.sub(r"\s+", "", text)


def strip_unit_suffix(value: object) -> str:
    return UNIT_SUFFIX_PATTERN.sub("", str(value).strip()).strip()


def pollutant_name_keys(value: object) -> set[str]:
    return {
        key
        for key in (
            normalize_name(value),
            normalize_name(strip_unit_suffix(value)),
        )
        if key
    }


def clean_pollutant_inputs(pollutants: list[str]) -> list[str]:
    cleaned: list[str] = []
    seen: set[str] = set()
    for pollutant in pollutants:
        name = str(pollutant).strip()
        if not name:
            continue
        identity = normalize_name(strip_unit_suffix(name))
        if identity not in seen:
            cleaned.append(name)
            seen.add(identity)
    if not cleaned:
        raise ValueError("At least one non-empty pollutant name is required")
    return cleaned


def validate_paths(input_path: Path, output_path: Path) -> None:
    if not input_path.is_file():
        raise FileNotFoundError(f"Input workbook not found: {input_path}")
    if input_path.suffix.lower() not in SUPPORTED_SUFFIXES:
        raise ValueError(f"Input workbook must be an .xlsx or .xlsm file: {input_path}")
    if output_path.suffix.lower() not in SUPPORTED_SUFFIXES:
        raise ValueError(
            f"Output workbook must be an .xlsx or .xlsm file: {output_path}"
        )
    if input_path == output_path:
        raise ValueError("Input and output paths must be different")
    if input_path.suffix.lower() != output_path.suffix.lower():
        raise ValueError("Input and output workbook extensions must match")


def filter_excel_pollutants(
    input_path: str | Path,
    pollutants: list[str],
    output_path: str | Path,
    *,
    allow_missing: bool = False,
) -> tuple[Path, list[SheetFilterSummary]]:
    input_file = resolve_path(input_path)
    output_file = resolve_path(output_path)
    validate_paths(input_file, output_file)
    requested = clean_pollutant_inputs(pollutants)
    requested_keys = {name: pollutant_name_keys(name) for name in requested}

    keep_vba = input_file.suffix.lower() == ".xlsm"
    workbook = load_workbook(input_file, keep_vba=keep_vba)
    sheet_plans: dict[str, set[int]] = {}
    matched = {name: False for name in requested}

    # Build and validate every sheet plan before modifying the workbook.
    for worksheet in workbook.worksheets:
        keep_columns: set[int] = set()
        has_time_column = False
        for column_index in range(1, worksheet.max_column + 1):
            header = worksheet.cell(row=1, column=column_index).value
            if normalize_name(header) == normalize_name(TIME_COLUMN):
                keep_columns.add(column_index)
                has_time_column = True
                continue

            header_keys = pollutant_name_keys(header)
            for name, name_keys in requested_keys.items():
                if header_keys & name_keys:
                    keep_columns.add(column_index)
                    matched[name] = True

        if not has_time_column:
            raise ValueError(
                f"Sheet '{worksheet.title}' does not contain the '{TIME_COLUMN}' "
                "column in row 1"
            )
        sheet_plans[worksheet.title] = keep_columns

    missing = [name for name, was_matched in matched.items() if not was_matched]
    if missing and not allow_missing:
        missing_text = ", ".join(missing)
        raise ValueError(
            f"Requested pollutants were not found in any worksheet: {missing_text}"
        )

    summaries: list[SheetFilterSummary] = []
    for worksheet in workbook.worksheets:
        original_columns = worksheet.max_column
        keep_columns = sheet_plans[worksheet.title]
        for column_index in range(original_columns, 0, -1):
            if column_index not in keep_columns:
                worksheet.delete_cols(column_index)
        summaries.append(
            SheetFilterSummary(
                sheet_name=worksheet.title,
                original_columns=original_columns,
                retained_columns=len(keep_columns),
            )
        )

    output_file.parent.mkdir(parents=True, exist_ok=True)
    temporary_file = output_file.with_name(
        f".{output_file.stem}.tmp{output_file.suffix}"
    )
    try:
        workbook.save(temporary_file)
        temporary_file.replace(output_file)
    finally:
        workbook.close()
        temporary_file.unlink(missing_ok=True)

    return output_file, summaries


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Retain only the time column and specified pollutant columns in every "
            "worksheet of a monitoring workbook."
        )
    )
    parser.add_argument(
        "input_path",
        nargs="?",
        default=str(INPUT_FILE_PATH),
        help="Input .xlsx or .xlsm workbook (defaults to INPUT_FILE_PATH)",
    )
    parser.add_argument(
        "--pollutants",
        nargs="+",
        default=POLLUTANTS,
        help="Pollutant names to retain (defaults to POLLUTANTS)",
    )
    parser.add_argument(
        "--output",
        default=str(OUTPUT_FILE_PATH),
        help="Output workbook path (defaults to OUTPUT_FILE_PATH)",
    )
    parser.add_argument(
        "--allow-missing",
        action="store_true",
        default=ALLOW_MISSING_POLLUTANTS,
        help="Continue when one or more requested pollutants are absent",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    output_path, summaries = filter_excel_pollutants(
        input_path=args.input_path,
        pollutants=args.pollutants,
        output_path=args.output,
        allow_missing=args.allow_missing,
    )

    print(f"Output workbook: {output_path}")
    for summary in summaries:
        print(
            f"  {summary.sheet_name}: "
            f"kept {summary.retained_columns}/{summary.original_columns} columns"
        )


if __name__ == "__main__":
    main()
